import shutil
import subprocess
import tempfile
from pathlib import Path

from app.core.config import settings
from app.core.exceptions import unprocessable
from app.utils.file_utils import clean_text


def _tesseract_options() -> tuple[object, str]:
    try:
        import pytesseract
    except ImportError as error:
        raise unprocessable("이미지 OCR 패키지가 설치되지 않았습니다. requirements.txt의 OCR 의존성을 설치해 주세요.") from error

    if settings.tesseract_cmd:
        executable = Path(settings.tesseract_cmd)
        if not executable.is_file():
            raise unprocessable("TESSERACT_CMD에 지정한 실행 파일을 찾을 수 없습니다. .env의 경로를 확인해 주세요.")
        pytesseract.pytesseract.tesseract_cmd = str(executable)

    config = ""
    if settings.tessdata_dir:
        data_directory = Path(settings.tessdata_dir)
        if not data_directory.is_dir():
            raise unprocessable("TESSDATA_DIR에 지정한 언어 데이터 폴더를 찾을 수 없습니다. .env의 경로를 확인해 주세요.")
        config = f'--tessdata-dir "{data_directory}"'
    return pytesseract, config


def _ocr_image(path: Path) -> str:
    try:
        from PIL import Image
        import pytesseract
    except ImportError as error:
        raise unprocessable("이미지 OCR 패키지가 설치되지 않았습니다. requirements.txt의 OCR 의존성을 설치해 주세요.") from error
    try:
        pytesseract, config = _tesseract_options()
        with Image.open(path) as image:
            return pytesseract.image_to_string(image, lang="kor+eng", config=config)
    except pytesseract.TesseractNotFoundError as error:
        raise unprocessable("Tesseract OCR 엔진 또는 한국어 언어 데이터(kor)가 설치되지 않았습니다.") from error
    except pytesseract.TesseractError as error:
        raise unprocessable("Tesseract가 한국어 언어 데이터(kor)를 읽을 수 없습니다. TESSDATA_DIR 경로와 kor.traineddata 파일을 확인해 주세요.") from error


def _ocr_pdf(path: Path) -> str:
    try:
        from pdf2image import convert_from_path
        from pdf2image.exceptions import PDFInfoNotInstalledError, PDFPageCountError, PDFSyntaxError
        import pytesseract
    except ImportError as error:
        raise unprocessable("스캔 PDF OCR 패키지가 설치되지 않았습니다. requirements.txt의 OCR 의존성을 설치해 주세요.") from error
    try:
        pytesseract, config = _tesseract_options()
        pages = convert_from_path(str(path))
        return "\n".join(pytesseract.image_to_string(page, lang="kor+eng", config=config) for page in pages)
    except pytesseract.TesseractNotFoundError as error:
        raise unprocessable("Tesseract OCR 엔진 또는 한국어 언어 데이터(kor)가 설치되지 않았습니다.") from error
    except pytesseract.TesseractError as error:
        raise unprocessable("Tesseract가 한국어 언어 데이터(kor)를 읽을 수 없습니다. TESSDATA_DIR 경로와 kor.traineddata 파일을 확인해 주세요.") from error
    except PDFInfoNotInstalledError as error:
        raise unprocessable("스캔 PDF OCR에 필요한 Poppler를 찾을 수 없습니다. Poppler의 bin 폴더를 PATH에 추가한 뒤 서버를 다시 시작해 주세요.") from error
    except (PDFPageCountError, PDFSyntaxError) as error:
        raise unprocessable("스캔 PDF의 페이지 정보를 읽을 수 없습니다. 암호화·손상 여부를 확인하거나 다른 PDF 뷰어에서 다시 저장한 파일을 업로드해 주세요.") from error
    except Exception as error:
        raise unprocessable("스캔 PDF OCR 처리 중 오류가 발생했습니다. 파일 상태와 OCR 실행 환경을 확인해 주세요.") from error


def _extract_legacy(path: Path) -> str:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise unprocessable("DOC, HWP, PPT 파일은 LibreOffice가 설치된 환경에서만 분석할 수 있습니다.")
    with tempfile.TemporaryDirectory() as output_dir:
        result = subprocess.run(
            [executable, "--headless", "--convert-to", "txt:Text", "--outdir", output_dir, str(path)],
            capture_output=True, text=True, timeout=60, check=False,
        )
        text_path = Path(output_dir) / f"{path.stem}.txt"
        if result.returncode != 0 or not text_path.exists():
            raise unprocessable("문서를 읽을 수 없습니다. LibreOffice 변환 결과와 파일 권한을 확인해 주세요.")
        return text_path.read_text(encoding="utf-8", errors="replace")


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".txt":
            try:
                raw = path.read_text(encoding="utf-8-sig")
            except UnicodeDecodeError:
                raw = path.read_text(encoding="cp949")
        elif suffix == ".pdf":
            from pypdf import PdfReader
            raw = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
            if not clean_text(raw):
                raw = _ocr_pdf(path)
        elif suffix == ".docx":
            from docx import Document
            document = Document(str(path))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            tables = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            raw = "\n".join(paragraphs + tables)
        elif suffix == ".pptx":
            from pptx import Presentation
            presentation = Presentation(str(path))
            raw = "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text"))
        elif suffix == ".hwpx":
            from hwpx import HwpxDocument
            raw = HwpxDocument.open(path).export_text()
        elif suffix in {".jpg", ".jpeg", ".png", ".tif", ".tiff"}:
            raw = _ocr_image(path)
        else:
            raw = _extract_legacy(path)
    except Exception as error:
        if hasattr(error, "status_code"):
            raise
        raise unprocessable("문서를 읽을 수 없습니다. 파일 상태 또는 접근 권한을 확인해 주세요.") from error

    text = clean_text(raw)
    if not text:
        raise unprocessable("분석할 내용이 없는 문서입니다. 텍스트가 포함된 문서를 업로드해 주세요.")
    if len(text) > settings.max_text_length:
        raise unprocessable("추출된 텍스트가 100,000,000자를 초과했습니다. 문서를 정리한 뒤 다시 업로드해 주세요.")
    return text
